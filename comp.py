import re
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import os as _os
_BASE = _os.path.dirname(_os.path.abspath(__file__))
FILE1 = _os.path.join(_BASE, '전.xlsx')
FILE2 = _os.path.join(_BASE, '후.xlsx')
SHEET = '공종별집계표'
OUT   = _os.path.join(_BASE, '공종별집계표_1대1비교(전후).xlsx')

wb1 = load_workbook(FILE1, data_only=True)
wb2 = load_workbook(FILE2, data_only=True)
ws1 = wb1[SHEET]
ws2 = wb2[SHEET]

def extract_items(ws):
    items = []; grand = None
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, values_only=True):
        raw = row[0]
        if not raw: continue
        name = str(raw).strip()
        if ('합' in name and '계' in name and '[' in name) or '합     계' in name:
            if isinstance(row[11], (int, float)):
                grand = {'mat': row[5] or 0, 'lab': row[7] or 0,
                         'exp': row[9] or 0, 'tot': row[11] or 0}
            continue
        if not isinstance(row[11], (int, float)): continue
        m = re.match(r'^(\d+)\s', name)
        level = len(m.group(1)) // 2 if m else 1
        clean = re.sub(r'^\d{2,6}\s+', '', name).strip()
        items.append({'name': clean, 'mat': row[5] or 0, 'lab': row[7] or 0,
                      'exp': row[9] or 0, 'tot': row[11] or 0, 'level': level})
    return items, grand

items1, grand1 = extract_items(ws1)
items2, grand2 = extract_items(ws2)

def norm(s): return re.sub(r'\s+', '', s).lower()

map1 = {norm(i['name']): i for i in items1}
result_rows = []; used1 = set()

for it2 in items2:
    key = norm(it2['name'])
    if key in map1:
        it1 = map1[key]; used1.add(key)
        it1['level'] = it2['level']
        result_rows.append(('match', it1, it2))
    else:
        result_rows.append(('add', None, it2))

for it1 in items1:
    if norm(it1['name']) not in used1:
        result_rows.append(('del', it1, None))

match_cnt = sum(1 for t,_,_ in result_rows if t=='match')
add_cnt   = sum(1 for t,_,_ in result_rows if t=='add')
del_cnt   = sum(1 for t,_,_ in result_rows if t=='del')
print(f'매칭:{match_cnt} / 내역추가:{add_cnt} / 내역삭제:{del_cnt}')
print(f'전.xlsx 합계: {grand1["tot"]:,}')
print(f'후.xlsx 합계: {grand2["tot"]:,}')
print(f'증감: {grand2["tot"]-grand1["tot"]:+,}')

# ── 엑셀 생성 ──
out = Workbook()
ws  = out.active
ws.title = '공종별집계표_비교'

thin = Side(style='thin', color='CCCCCC')
med  = Side(style='medium', color='888888')
def bd(c, heavy=False):
    s = med if heavy else thin
    c.border = Border(left=s, right=s, top=s, bottom=s)

HDR='1F3864'; WHT='FFFFFF'
col_widths = [3, 36, 16, 16, 14, 16, 16, 16, 14, 16, 16, 10]
col_names  = ['L','항목명','① 재료비','① 노무비','① 경비','① 합계',
              '② 재료비','② 노무비','② 경비','② 합계','증감액(합계)','비고']
for i,(h,w) in enumerate(zip(col_names,col_widths),1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.merge_cells('A1:L1')
c = ws.cell(1,1,'공종별집계표 1대1 비교   ① 전.xlsx  vs  ② 후.xlsx   【원 단위】')
c.font = Font(name='맑은 고딕',bold=True,size=12,color=WHT)
c.fill = PatternFill('solid',start_color=HDR)
c.alignment = Alignment(horizontal='center',vertical='center')
ws.row_dimensions[1].height = 26

for col,h in enumerate(col_names,1):
    c = ws.cell(2,col,h)
    c.font = Font(name='맑은 고딕',bold=True,size=9,color=WHT)
    c.fill = PatternFill('solid',start_color='2E75B6')
    c.alignment = Alignment(horizontal='center',vertical='center',wrap_text=True)
    bd(c,True)
ws.row_dimensions[2].height = 30
ws.freeze_panes = 'A3'

def fnum(c,v,bg,bold=False):
    c.value = v if v!=0 else None
    c.number_format = '#,##0;(#,##0);-'
    c.font = Font(name='맑은 고딕',size=9,bold=bold,
                  color='C00000' if isinstance(v,(int,float)) and v<0 else '000000')
    c.fill = PatternFill('solid',start_color=bg)
    c.alignment = Alignment(horizontal='right',vertical='center')
    bd(c)

def ftxt(c,v,bg,bold=False,center=False):
    c.value = v
    c.font = Font(name='맑은 고딕',size=9,bold=bold)
    c.fill = PatternFill('solid',start_color=bg)
    c.alignment = Alignment(horizontal='center' if center else 'left',vertical='center',wrap_text=True)
    bd(c)

r = 3
for rtype,it1,it2 in result_rows:
    lv = (it2['level'] if it2 else it1.get('level',2)) or 1
    indent = '  '*(lv-1)
    bold = lv<=2

    if rtype=='match':
        diff = it2['tot']-it1['tot']
        bg = 'E2EFDA' if diff>0 else ('FCE4D6' if diff<0 else 'F5F5F5')
        ftxt(ws.cell(r,1),str(lv),bg,center=True)
        ftxt(ws.cell(r,2),indent+it1['name'],bg,bold=bold)
        fnum(ws.cell(r,3),it1['mat'],bg)
        fnum(ws.cell(r,4),it1['lab'],bg)
        fnum(ws.cell(r,5),it1['exp'],bg)
        fnum(ws.cell(r,6),it1['tot'],bg,bold=bold)
        fnum(ws.cell(r,7),it2['mat'],bg)
        fnum(ws.cell(r,8),it2['lab'],bg)
        fnum(ws.cell(r,9),it2['exp'],bg)
        fnum(ws.cell(r,10),it2['tot'],bg,bold=bold)
        fnum(ws.cell(r,11),diff,bg,bold=bold)
        pct = diff/it1['tot'] if it1['tot'] else 0
        c2 = ws.cell(r,12,f'{pct:+.1%}' if it1['tot'] else '-')
        c2.font=Font(name='맑은 고딕',size=9)
        c2.fill=PatternFill('solid',start_color=bg)
        c2.alignment=Alignment(horizontal='center',vertical='center')
        bd(c2)

    elif rtype=='add':
        bg='DDEBF7'
        ftxt(ws.cell(r,1),str(lv),bg,center=True)
        ftxt(ws.cell(r,2),indent+it2['name'],bg,bold=bold)
        for ci in [3,4,5,6]: ftxt(ws.cell(r,ci),'-',bg,center=True)
        fnum(ws.cell(r,7),it2['mat'],bg)
        fnum(ws.cell(r,8),it2['lab'],bg)
        fnum(ws.cell(r,9),it2['exp'],bg)
        fnum(ws.cell(r,10),it2['tot'],bg,bold=True)
        fnum(ws.cell(r,11),it2['tot'],bg,bold=True)
        c2=ws.cell(r,12,'내역추가')
        c2.font=Font(name='맑은 고딕',size=9,bold=True,color=WHT)
        c2.fill=PatternFill('solid',start_color='2E75B6')
        c2.alignment=Alignment(horizontal='center',vertical='center')
        bd(c2)

    else:
        bg='FFF2CC'
        ftxt(ws.cell(r,1),str(lv),bg,center=True)
        ftxt(ws.cell(r,2),indent+it1['name'],bg,bold=bold)
        fnum(ws.cell(r,3),it1['mat'],bg)
        fnum(ws.cell(r,4),it1['lab'],bg)
        fnum(ws.cell(r,5),it1['exp'],bg)
        fnum(ws.cell(r,6),it1['tot'],bg,bold=True)
        for ci in [7,8,9,10]: ftxt(ws.cell(r,ci),'-',bg,center=True)
        fnum(ws.cell(r,11),-it1['tot'],bg,bold=True)
        c2=ws.cell(r,12,'내역삭제')
        c2.font=Font(name='맑은 고딕',size=9,bold=True,color=WHT)
        c2.fill=PatternFill('solid',start_color='C55A11')
        c2.alignment=Alignment(horizontal='center',vertical='center')
        bd(c2)

    ws.row_dimensions[r].height = 17
    r += 1

# 합계 행
r += 1
ws.merge_cells(f'A{r}:B{r}')
c = ws.cell(r,1,'■ 공식 합계 ([합 계] 행 기준)')
c.font=Font(name='맑은 고딕',bold=True,size=10,color=WHT)
c.fill=PatternFill('solid',start_color=HDR)
c.alignment=Alignment(horizontal='center',vertical='center')
ws.cell(r,2).fill=PatternFill('solid',start_color=HDR)
bg='D6E4FF'
fnum(ws.cell(r,3), grand1['mat'],bg,True)
fnum(ws.cell(r,4), grand1['lab'],bg,True)
fnum(ws.cell(r,5), grand1['exp'],bg,True)
fnum(ws.cell(r,6), grand1['tot'],bg,True)
fnum(ws.cell(r,7), grand2['mat'],bg,True)
fnum(ws.cell(r,8), grand2['lab'],bg,True)
fnum(ws.cell(r,9), grand2['exp'],bg,True)
fnum(ws.cell(r,10),grand2['tot'],bg,True)
diff_tot=grand2['tot']-grand1['tot']
fnum(ws.cell(r,11),diff_tot,bg,True)
pct=diff_tot/grand1['tot'] if grand1['tot'] else 0
c2=ws.cell(r,12,f'{pct:+.1%}')
c2.font=Font(name='맑은 고딕',size=10,bold=True)
c2.fill=PatternFill('solid',start_color=bg)
c2.alignment=Alignment(horizontal='center',vertical='center')
bd(c2)
ws.row_dimensions[r].height=22

out.save(OUT)
print('저장완료:', OUT)
