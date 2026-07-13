import re, io, os
from flask import Flask, request, jsonify, send_file, render_template
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 100 * 1024 * 1024

UPLOAD = os.path.join(os.path.dirname(__file__), 'uploads')
os.makedirs(UPLOAD, exist_ok=True)

def norm(s):
    return re.sub(r'\s+', '', str(s)).lower()

def norm_name(s):
    return re.sub(r'[\s\-_\.\(\)\[\]①②③④⑤]', '', str(s)).lower()

def detect_amount_cols(ws):
    """재료비/노무비/경비/합계 컬럼 자동 탐지
    전체 스캔 후 우선순위: A(금액4개) > C(금액3개) > B(재료/노무) > D(금액1개) > E(합계만)
    """
    cand_A = cand_B = cand_C = None
    single_d = tot_e = None

    pa = get_print_area_bounds(ws)
    hdr_start = pa[0] if pa else 1
    hdr_end   = min(hdr_start + 11, pa[1] if pa else ws.max_row)
    for row in ws.iter_rows(min_row=hdr_start, max_row=hdr_end, values_only=True):
        n = [re.sub(r'\s+', '', str(c)).lower() if c else '' for c in row]
        gumack = [i for i, v in enumerate(n) if '금액' in v]

        if len(gumack) >= 4 and cand_A is None:
            cand_A = (gumack[0], gumack[1], gumack[2], gumack[3])

        if len(gumack) >= 3 and cand_C is None:
            htot = next((i for i, v in enumerate(n) if '합계' in v and i not in gumack), None)
            cand_C = (gumack[0], gumack[1], gumack[2], htot if htot else max(gumack) + 1)

        mat = next((i for i, v in enumerate(n) if '재료' in v), None)
        lab = next((i for i, v in enumerate(n) if '노무' in v), None)
        exp = next((i for i, v in enumerate(n) if '경비' in v), None)
        tot = next((i for i, v in enumerate(n) if i > 0 and '합계' in v and '소' not in v), None)
        if mat is not None and lab is not None and tot is not None and cand_B is None:
            cand_B = (mat, lab, exp if exp is not None else lab + 1, tot)

        if len(gumack) == 1 and gumack[0] >= 4 and single_d is None:
            single_d = gumack[0]
        if tot is not None and mat is None and lab is None and tot_e is None:
            tot_e = tot

    if cand_A: return cand_A
    if cand_C: return cand_C
    if cand_B: return cand_B
    if single_d:
        ci = single_d; return max(0, ci-3), max(0, ci-2), max(0, ci-1), ci
    if tot_e:
        ci = tot_e; return max(0, ci-3), max(0, ci-2), max(0, ci-1), ci
    return 5, 7, 9, 11

def get_print_area_bounds(ws):
    """인쇄 영역의 (min_row, max_row) 반환. 없으면 None."""
    pa = ws.print_area
    if not pa:
        return None
    area = pa.split(',')[0].strip()
    if '!' in area:
        area = area.split('!')[1]
    area = area.replace('$', '')
    try:
        from openpyxl.utils import range_boundaries
        min_col, min_row, max_col, max_row = range_boundaries(area)
        return min_row, max_row
    except Exception:
        return None

def extract_items(ws):
    ci_mat, ci_lab, ci_exp, ci_tot = detect_amount_cols(ws)
    items = []; grand = None

    # 인쇄 영역이 있으면 해당 범위만, 없으면 전체 시트
    pa = get_print_area_bounds(ws)
    scan_start = pa[0] if pa else 1
    scan_end   = pa[1] if pa else ws.max_row

    # 데이터 시작행 자동탐지 (인쇄 영역 내에서)
    data_start = scan_start + 4
    for ri, row in enumerate(ws.iter_rows(min_row=scan_start,
                                           max_row=min(scan_start + 11, scan_end),
                                           values_only=True), scan_start):
        num_count = sum(1 for c in row if isinstance(c, (int, float)) and c != 0)
        if num_count >= 2:
            data_start = ri; break

    for row in ws.iter_rows(min_row=data_start, max_row=scan_end, values_only=True):
        if not row: continue
        # 이름: 첫 번째 비어있지 않은 텍스트 셀
        raw = None
        for cell in row:
            if cell and isinstance(cell, str) and len(str(cell).strip()) > 0:
                raw = cell; break
        if raw is None: continue
        name = str(raw).strip()

        # 합계행
        nc = re.sub(r'\s+', '', name)
        if '합' in nc and '계' in nc:
            tot_val = row[ci_tot] if ci_tot < len(row) else None
            if isinstance(tot_val, (int, float)) and abs(tot_val) > 0:
                grand = {
                    'mat': (row[ci_mat] if ci_mat < len(row) else None) or 0,
                    'lab': (row[ci_lab] if ci_lab < len(row) else None) or 0,
                    'exp': (row[ci_exp] if ci_exp < len(row) else None) or 0,
                    'tot': tot_val,
                }
            continue

        tot_val = row[ci_tot] if ci_tot < len(row) else None
        if not isinstance(tot_val, (int, float)) or tot_val == 0: continue

        # 코드 추출: 숫자 접두 or col[13]
        m = re.match(r'^(\d{4,})\s+', name)
        if m:
            code = m.group(1)
            level = max(1, len(code) // 2)
            display_name = name[m.end():].strip()
        else:
            code_cell = row[13] if len(row) > 13 else None
            code = str(code_cell).strip() if code_cell and str(code_cell).strip() not in ('', 'None') else None
            level = 2
            display_name = name

        def _n(v): return v if isinstance(v, (int, float)) else 0
        items.append({
            'name': display_name,
            'raw_name': name,
            'code': code,
            'mat': _n(row[ci_mat] if ci_mat < len(row) else None),
            'lab': _n(row[ci_lab] if ci_lab < len(row) else None),
            'exp': _n(row[ci_exp] if ci_exp < len(row) else None),
            'tot': tot_val,
            'level': level,
        })
    return items, grand


def _fuzzy_match(name1, name2):
    n1 = norm_name(name1); n2 = norm_name(name2)
    if not n1 or not n2: return False
    short, long_ = (n1, n2) if len(n1) <= len(n2) else (n2, n1)
    return len(short) >= 4 and short in long_


def do_compare(ws1, ws2):
    items1, grand1 = extract_items(ws1)
    items2, grand2 = extract_items(ws2)

    # 전략 판단
    num_codes1 = {i['code'] for i in items1 if i['code'] and re.match(r'^\d{4,}$', str(i['code']))}
    num_codes2 = {i['code'] for i in items2 if i['code'] and re.match(r'^\d{4,}$', str(i['code']))}
    use_numcode = bool(num_codes1 & num_codes2)

    str_codes1 = {i['code'] for i in items1 if i['code'] and not re.match(r'^\d{4,}$', str(i['code']))}
    str_codes2 = {i['code'] for i in items2 if i['code'] and not re.match(r'^\d{4,}$', str(i['code']))}
    use_strcode = bool(str_codes1 & str_codes2)

    code_map1 = {}
    if use_numcode:
        code_map1 = {i['code']: i for i in items1 if i['code'] and re.match(r'^\d{4,}$', str(i['code']))}
    elif use_strcode:
        code_map1 = {i['code']: i for i in items1 if i['code']}

    norm_map1 = {}
    for i in items1:
        k = norm(i['name'])
        if k not in norm_map1: norm_map1[k] = i
        rk = norm(i['raw_name'])
        if rk not in norm_map1: norm_map1[rk] = i

    used1 = set()  # item id (id(obj))

    def find_match(it2):
        if it2['code'] and it2['code'] in code_map1:
            it1 = code_map1[it2['code']]
            if id(it1) not in used1:
                used1.add(id(it1)); return it1
        for k in [norm(it2['name']), norm(it2['raw_name'])]:
            if k and k in norm_map1:
                it1 = norm_map1[k]
                if id(it1) not in used1:
                    used1.add(id(it1)); return it1
        for it1 in items1:
            if id(it1) in used1: continue
            if _fuzzy_match(it2['name'], it1['name']):
                used1.add(id(it1)); return it1
        return None

    # 1단계: items2 전체에 대해 매칭 사전 계산
    matches2 = []  # [(idx2, it2, matched_it1_or_None), ...]
    for idx2, it2 in enumerate(items2):
        it1 = find_match(it2)
        matches2.append((idx2, it2, it1))

    # items1 id → items2 index 역매핑
    id1_to_idx2 = {id(it1): idx2 for idx2, it2, it1 in matches2 if it1 is not None}

    # 2단계: items1 순서를 기준으로 merge
    # - matched items1 → 앞에 끼어있는 add(items2 전용) 항목 먼저 출력 → match 출력
    # - unmatched items1 → del로 해당 위치에 출력
    # 이렇게 하면 양쪽 시트의 원래 순서가 최대한 보존됨
    result = []
    next_idx2 = 0  # items2에서 다음으로 처리할 인덱스

    for it1 in items1:
        if id(it1) in id1_to_idx2:
            idx2 = id1_to_idx2[id(it1)]
            it2 = items2[idx2]
            # idx2 이전에 있는 미매칭 items2 항목(add)을 먼저 출력
            while next_idx2 < idx2:
                _, it2_pre, it1_pre = matches2[next_idx2]
                if it1_pre is None:
                    result.append(('add', None, it2_pre))
                next_idx2 += 1
            next_idx2 = idx2 + 1
            it1_copy = dict(it1); it1_copy['level'] = it2['level']
            result.append(('match', it1_copy, it2))
        else:
            # 삭제 항목: items1의 원래 위치에 출력
            result.append(('del', it1, None))

    # 남은 items2 add 항목 후미 출력
    while next_idx2 < len(matches2):
        _, it2_rem, it1_rem = matches2[next_idx2]
        if it1_rem is None:
            result.append(('add', None, it2_rem))
        next_idx2 += 1

    return result, grand1, grand2

def build_excel(rows, grand1, grand2, label1, label2, sheet_name):
    out = Workbook(); ws = out.active
    ws.title = f'{sheet_name}_비교'[:31]
    thin = Side(style='thin', color='CCCCCC')
    med  = Side(style='medium', color='888888')
    def bd(c, heavy=False):
        s = med if heavy else thin
        c.border = Border(left=s, right=s, top=s, bottom=s)

    HDR='1F3864'; WHT='FFFFFF'
    widths = [3, 36, 14, 14, 12, 14, 14, 14, 12, 14, 16, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.merge_cells('A1:L1')
    c = ws.cell(1, 1, f'{sheet_name} 1대1 비교   ① {label1}  vs  ② {label2}   【원 단위】')
    c.font = Font(name='맑은 고딕', bold=True, size=12, color=WHT)
    c.fill = PatternFill('solid', start_color=HDR)
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 26

    ws.merge_cells('A2:B2'); ws.merge_cells('C2:F2'); ws.merge_cells('G2:J2')
    for ci, col_name, bg in [(3, f'① {label1}', '1F3864'),
                              (7, f'② {label2}', 'C55A11'),
                              (11, '증감액', '1A5C38'),
                              (12, '비고', '374151')]:
        c = ws.cell(2, ci, col_name)
        c.font = Font(name='맑은 고딕', bold=True, size=10, color=WHT)
        c.fill = PatternFill('solid', start_color=bg)
        c.alignment = Alignment(horizontal='center', vertical='center')
        bd(c, True)
    ws.cell(2, 1).fill = PatternFill('solid', start_color='374151')
    ws.row_dimensions[2].height = 22

    sub = ['L', '', '재료비', '노무비', '경비', '합계', '재료비', '노무비', '경비', '합계', '(합계)', '']
    sub_bgs = ['374151','374151','2C4F84','2C4F84','2C4F84','2C4F84','A04A0E','A04A0E','A04A0E','A04A0E','1F6B3E','374151']
    for ci, (h, bg) in enumerate(zip(sub, sub_bgs), 1):
        c = ws.cell(3, ci, h)
        c.font = Font(name='맑은 고딕', bold=True, size=8, color=WHT)
        c.fill = PatternFill('solid', start_color=bg)
        c.alignment = Alignment(horizontal='center', vertical='center')
        bd(c, True)
    ws.row_dimensions[3].height = 22
    ws.freeze_panes = 'A4'

    def fnum(c, v, bg, bold=False):
        c.value = v if v != 0 else None
        c.number_format = '#,##0;(#,##0);-'
        c.font = Font(name='맑은 고딕', size=9, bold=bold,
                      color='C00000' if isinstance(v, (int, float)) and v < 0 else '000000')
        c.fill = PatternFill('solid', start_color=bg)
        c.alignment = Alignment(horizontal='right', vertical='center')
        bd(c)
    def ftxt(c, v, bg, bold=False, center=False):
        c.value = v
        c.font = Font(name='맑은 고딕', size=9, bold=bold)
        c.fill = PatternFill('solid', start_color=bg)
        c.alignment = Alignment(horizontal='center' if center else 'left',
                                vertical='center', wrap_text=True)
        bd(c)

    r = 4
    for rtype, it1, it2 in rows:
        lv = (it2['level'] if it2 else (it1 or {}).get('level', 2)) or 2
        indent = '  ' * (min(lv, 3) - 1); bold = lv <= 2
        if rtype == 'match':
            diff = it2['tot'] - it1['tot']
            bg = 'E2EFDA' if diff > 0 else ('FCE4D6' if diff < 0 else 'F5F5F5')
            ftxt(ws.cell(r,1), str(lv), bg, center=True)
            ftxt(ws.cell(r,2), indent+it1['name'], bg, bold=bold)
            fnum(ws.cell(r,3), it1['mat'], bg); fnum(ws.cell(r,4), it1['lab'], bg)
            fnum(ws.cell(r,5), it1['exp'], bg); fnum(ws.cell(r,6), it1['tot'], bg, bold=bold)
            fnum(ws.cell(r,7), it2['mat'], bg); fnum(ws.cell(r,8), it2['lab'], bg)
            fnum(ws.cell(r,9), it2['exp'], bg); fnum(ws.cell(r,10), it2['tot'], bg, bold=bold)
            fnum(ws.cell(r,11), diff, bg, bold=bold)
            pct = diff / it1['tot'] if it1['tot'] else 0
            c2 = ws.cell(r, 12, f'{pct:+.1%}' if it1['tot'] else '-')
            c2.font = Font(name='맑은 고딕', size=9)
            c2.fill = PatternFill('solid', start_color=bg)
            c2.alignment = Alignment(horizontal='center', vertical='center'); bd(c2)
        elif rtype == 'add':
            bg = 'DDEBF7'
            ftxt(ws.cell(r,1), str(lv), bg, center=True)
            ftxt(ws.cell(r,2), indent+it2['name'], bg, bold=bold)
            for ci in [3,4,5,6]: ftxt(ws.cell(r,ci), '-', bg, center=True)
            fnum(ws.cell(r,7), it2['mat'], bg); fnum(ws.cell(r,8), it2['lab'], bg)
            fnum(ws.cell(r,9), it2['exp'], bg); fnum(ws.cell(r,10), it2['tot'], bg, bold=True)
            fnum(ws.cell(r,11), it2['tot'], bg, bold=True)
            c2 = ws.cell(r, 12, '내역추가')
            c2.font = Font(name='맑은 고딕', size=9, bold=True, color=WHT)
            c2.fill = PatternFill('solid', start_color='2E75B6')
            c2.alignment = Alignment(horizontal='center', vertical='center'); bd(c2)
        else:
            bg = 'FFF2CC'
            ftxt(ws.cell(r,1), str(lv), bg, center=True)
            ftxt(ws.cell(r,2), indent+it1['name'], bg, bold=bold)
            fnum(ws.cell(r,3), it1['mat'], bg); fnum(ws.cell(r,4), it1['lab'], bg)
            fnum(ws.cell(r,5), it1['exp'], bg); fnum(ws.cell(r,6), it1['tot'], bg, bold=True)
            for ci in [7,8,9,10]: ftxt(ws.cell(r,ci), '-', bg, center=True)
            fnum(ws.cell(r,11), -it1['tot'], bg, bold=True)
            c2 = ws.cell(r, 12, '내역삭제')
            c2.font = Font(name='맑은 고딕', size=9, bold=True, color=WHT)
            c2.fill = PatternFill('solid', start_color='C55A11')
            c2.alignment = Alignment(horizontal='center', vertical='center'); bd(c2)
        ws.row_dimensions[r].height = 16; r += 1

    r += 1
    ws.merge_cells(f'A{r}:B{r}')
    c = ws.cell(r, 1, '■ 공식 합계 ([합 계] 행 기준)')
    c.font = Font(name='맑은 고딕', bold=True, size=10, color=WHT)
    c.fill = PatternFill('solid', start_color=HDR)
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(r, 2).fill = PatternFill('solid', start_color=HDR)
    bg = 'D6E4FF'
    g1 = grand1 or {'mat':0,'lab':0,'exp':0,'tot':0}
    g2 = grand2 or {'mat':0,'lab':0,'exp':0,'tot':0}
    fnum(ws.cell(r,3), g1['mat'], bg, True); fnum(ws.cell(r,4), g1['lab'], bg, True)
    fnum(ws.cell(r,5), g1['exp'], bg, True); fnum(ws.cell(r,6), g1['tot'], bg, True)
    fnum(ws.cell(r,7), g2['mat'], bg, True); fnum(ws.cell(r,8), g2['lab'], bg, True)
    fnum(ws.cell(r,9), g2['exp'], bg, True); fnum(ws.cell(r,10), g2['tot'], bg, True)
    diff_tot = g2['tot'] - g1['tot']
    fnum(ws.cell(r,11), diff_tot, bg, True)
    pct = diff_tot / g1['tot'] if g1['tot'] else 0
    c2 = ws.cell(r, 12, f'{pct:+.1%}')
    c2.font = Font(name='맑은 고딕', size=10, bold=True)
    c2.fill = PatternFill('solid', start_color=bg)
    c2.alignment = Alignment(horizontal='center', vertical='center'); bd(c2)
    ws.row_dimensions[r].height = 22

    buf = io.BytesIO(); out.save(buf); buf.seek(0)
    return buf


BASE_DIR = os.path.dirname(os.path.dirname(__file__))  # 내역서검토/

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/dev_preview', methods=['GET'])
def dev_preview():
    """개발용: 전.xlsx / 후.xlsx 자동 비교"""
    p1 = os.path.join(BASE_DIR, '전.xlsx')
    p2 = os.path.join(BASE_DIR, '후.xlsx')
    if not os.path.exists(p1) or not os.path.exists(p2):
        return jsonify(error='전.xlsx / 후.xlsx 파일이 없습니다'), 404
    wb1 = load_workbook(p1, data_only=True)
    wb2 = load_workbook(p2, data_only=True)
    sheet = '공종별집계표'
    rows, grand1, grand2 = do_compare(wb1[sheet], wb2[sheet])
    table = []
    for rtype, it1, it2 in rows:
        lv = (it2['level'] if it2 else (it1 or {}).get('level', 2)) or 2
        name = it2['name'] if it2 else it1['name']
        diff = (it2['tot'] - it1['tot']) if rtype == 'match' else \
               (it2['tot'] if rtype == 'add' else -it1['tot'])
        pct = diff / it1['tot'] if rtype == 'match' and it1['tot'] else None
        table.append({
            'type': rtype, 'level': lv, 'name': name,
            'mat1': it1['mat'] if it1 else None, 'lab1': it1['lab'] if it1 else None,
            'exp1': it1['exp'] if it1 else None, 'tot1': it1['tot'] if it1 else None,
            'mat2': it2['mat'] if it2 else None, 'lab2': it2['lab'] if it2 else None,
            'exp2': it2['exp'] if it2 else None, 'tot2': it2['tot'] if it2 else None,
            'diff': diff, 'pct': pct,
        })
    return jsonify({
        'table': table, 'grand1': grand1, 'grand2': grand2,
        'summary': {
            'match': sum(1 for r in rows if r[0]=='match'),
            'add':   sum(1 for r in rows if r[0]=='add'),
            'del':   sum(1 for r in rows if r[0]=='del'),
        },
        'label1': '전.xlsx', 'label2': '후.xlsx',
    })

@app.route('/api/sheets', methods=['POST'])
def get_sheets():
    files = {}
    for key in ['file1', 'file2']:
        f = request.files.get(key)
        if not f: return jsonify(error=f'{key} 없음'), 400
        path = os.path.join(UPLOAD, f'{key}_{f.filename}')
        f.save(path)
        wb = load_workbook(path, data_only=True, read_only=True)
        files[key] = {'name': f.filename, 'sheets': list(wb.sheetnames), 'path': path}
        wb.close()
    return jsonify(files)

@app.route('/api/compare', methods=['POST'])
def compare():
    data = request.json
    wb1 = load_workbook(data['path1'], data_only=True)
    wb2 = load_workbook(data['path2'], data_only=True)
    rows, grand1, grand2 = do_compare(wb1[data['sheet1']], wb2[data['sheet2']])

    table = []
    for rtype, it1, it2 in rows:
        lv = (it2['level'] if it2 else (it1 or {}).get('level', 2)) or 2
        name = it2['name'] if it2 else it1['name']
        diff = (it2['tot'] - it1['tot']) if rtype == 'match' else \
               (it2['tot'] if rtype == 'add' else -it1['tot'])
        pct = diff / it1['tot'] if rtype == 'match' and it1['tot'] else None
        table.append({
            'type': rtype, 'level': lv, 'name': name,
            'mat1': it1['mat'] if it1 else None, 'lab1': it1['lab'] if it1 else None,
            'exp1': it1['exp'] if it1 else None, 'tot1': it1['tot'] if it1 else None,
            'mat2': it2['mat'] if it2 else None, 'lab2': it2['lab'] if it2 else None,
            'exp2': it2['exp'] if it2 else None, 'tot2': it2['tot'] if it2 else None,
            'diff': diff, 'pct': pct,
        })

    return jsonify({
        'table': table, 'grand1': grand1, 'grand2': grand2,
        'summary': {
            'match': sum(1 for r in rows if r[0] == 'match'),
            'add':   sum(1 for r in rows if r[0] == 'add'),
            'del':   sum(1 for r in rows if r[0] == 'del'),
        },
        'path1': data['path1'], 'path2': data['path2'],
        'sheet1': data['sheet1'], 'sheet2': data['sheet2'],
        'label1': data['label1'], 'label2': data['label2'],
    })

@app.route('/api/download', methods=['POST'])
def download():
    data = request.json
    wb1 = load_workbook(data['path1'], data_only=True)
    wb2 = load_workbook(data['path2'], data_only=True)
    rows, grand1, grand2 = do_compare(wb1[data['sheet1']], wb2[data['sheet2']])
    buf = build_excel(rows, grand1, grand2, data['label1'], data['label2'], data['sheet1'])
    return send_file(buf, as_attachment=True,
                     download_name=f"{data['sheet1']}_1대1비교.xlsx",
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

if __name__ == '__main__':
    app.run(debug=True, port=5000)
